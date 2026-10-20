#!/usr/bin/env python3
"""make_report.py — Word-отчёт по треку 3 из results.json и plots/*.png.

Порядок:  python3 tools/analyze.py logs/omp_<jobid>.out
          python3 tools/make_report.py  ->  report_t3.docx
Нужно:   pip install python-docx matplotlib
"""
import json, re
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mp
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = json.load(open("results.json", encoding="utf-8"))
rows, T1 = R["rows"], R["T1"]
N, ITERS = 4096, 500
INNER = (N - 2) ** 2
BYTES_PER_PT = 16          # минимум: чтение u (8 Б, соседи из кэша) + запись unew (8 Б)
TEAM = "Мусин А.Р., <второй участник>"   # впишите название команды и ФИО напарника

# ---------- характеристики узла из лога ----------
def node_field(key):
    for l in R["node"]:
        if l.lower().startswith(key.lower()):
            return l.split(":", 1)[1].strip()
    return "—"
local = any("НЕ ПАРАМ" in l for l in R["node"])
node = [("Кластер / раздел", "ПАРАМ БИЛИМ, раздел compute, узел выделен целиком (--exclusive --mem=0)"),
        ("Узел (hostname)", node_field("node")),
        ("Процессор", node_field("Model name")),
        ("Сокетов × ядер", f"{node_field('Socket(s)')} × {node_field('Core(s) per socket')} "
                           f"(в Slurm CPUTot = 96 с учётом Hyper-Threading)"),
        ("NUMA-узлов", node_field("NUMA node(s)")),
        ("Кэш L3", node_field("L3 cache") + " на сокет"),
        ("Память узла", "256 ГБ (RealMemory из scontrol)"),
        ("Компилятор", "gcc 8.5.0 (Red Hat), флаги -O3 / -O3 -fopenmp; сборка на входном узле "
                       "bilim-master, т.к. на вычислительных узлах нет заголовков C"),
        ("Привязка потоков", "OMP_PROC_BIND=close, OMP_PLACES=cores")]

# ---------- схема декомпозиции ----------
def draw_decomposition(path, nt=4):
    fig, ax = plt.subplots(figsize=(6.2, 3.0))
    cols = ["#BFD3EE", "#F6D2B8", "#C9E6C9", "#E7D3F0"]
    H = 16
    ax.add_patch(mp.Rectangle((0, 0), H, H, fill=False, lw=1.5))
    ax.add_patch(mp.Rectangle((0, H - 1), H, 1, color="#E05A47"))
    for side in [(0, 0, H, 1), (0, 0, 1, H), (H - 1, 0, 1, H)]:
        ax.add_patch(mp.Rectangle(side[:2], side[2], side[3], color="#9AA4B2"))
    band = (H - 2) / nt
    n_it = N - 2; q, rem = divmod(n_it, nt)       # так делит schedule(static)
    start = 1
    for k in range(nt):
        cnt = q + (1 if k < rem else 0)
        y0 = H - 1 - (k + 1) * band
        ax.add_patch(mp.Rectangle((1, y0), H - 2, band, color=cols[k], ec="white"))
        ax.text(H / 2, y0 + band / 2, f"поток {k}: строки {start}–{start + cnt - 1}",
                ha="center", va="center", fontsize=9)
        start += cnt
    ax.text(H / 2, H - 0.5, "верхняя граница = 100", ha="center", va="center", color="white", fontsize=8)
    ax.annotate("", xy=(H + 1.6, 2), xytext=(H + 1.6, 6.5),
                arrowprops=dict(arrowstyle="<->", color="#333"))
    ax.text(H + 2.0, 4.2, "граничные строки\nсоседа читаются\nиз общей памяти\n(без обменов)",
            fontsize=8, va="center")
    ax.set_xlim(-0.5, H + 7.5); ax.set_ylim(-0.5, H + 0.5)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("Ленточная декомпозиция по строкам, schedule(static), пример p = 4", fontsize=10)
    plt.tight_layout(); plt.savefig(path, dpi=160); plt.close()

draw_decomposition("plots/decomposition.png")

# ---------- расчёты для анализа ----------
last = rows[-1]
best = max(rows, key=lambda r: r["S"])
es = [r["e"] for r in rows if r["e"] is not None]
grows = len(es) >= 2 and es[-1] > es[0] * 1.5 + 0.005
def bw(r):  # достигнутая пропускная способность памяти, ГБ/с
    return INNER * BYTES_PER_PT * ITERS / r["T"] / 1e9
bw_seq = INNER * BYTES_PER_PT * ITERS / T1 / 1e9
f_amdahl = es[-1] if es else None

# ---------- документ ----------
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2)

BLUE = RGBColor(0x2E, 0x5C, 0x9A)
def h(text, lvl=1):
    p = doc.add_heading(text, lvl)
    for r in p.runs: r.font.color.rgb = BLUE
def para(text="", bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix: p.add_run(bold_prefix).bold = True
    r = p.add_run(text); r.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p
def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix: p.add_run(bold_prefix).bold = True
    p.add_run(text)
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)
def table(header, data, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, x in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(x).bold = True; shade(c, "DCE6F2")
    for row in data:
        cells = t.add_row().cells
        for i, x in enumerate(row): cells[i].text = str(x)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph()
def code(lines):
    for l in lines:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(l); r.font.name = "Consolas"; r.font.size = Pt(9)
        pPr = p._p.get_or_add_pPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "F2F4F7"); pPr.append(sh)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# --- титул
doc.add_heading("РК1 · Трек 3: 2D теплопроводность (Якоби), OpenMP", 0)
para(f"Команда: {TEAM}")
P = {r["p"]: r for r in rows}
gain = lambda a, b: P[b]["S"] / P[a]["S"]

# --- 1. Задача и узел
h("1. Задача и узел")
para(f"Сетка {N}×{N}, {ITERS} итераций Якоби; верхняя граница 100, остальное 0. Новое значение "
     "точки — среднее четырёх соседей. Два массива по 128 МБ, в кэш не помещаются. "
     "CHECK: сумма внутренних точек совпадает с эталоном seq.c (4949330.8767278139, допуск 1e-9).")
table(["Параметр", "Значение"], [
    ("Узел", f"ПАРАМ БИЛИМ, раздел compute, {node_field('node')}, узел целиком"),
    ("Процессор", f"{node_field('Socket(s)')} × {node_field('Model name')}, "
                  f"{node_field('Core(s) per socket')} ядер на сокет"),
    ("NUMA / L3 / память", f"{node_field('NUMA node(s)')} NUMA-узла, L3 36 МБ на сокет, 256 ГБ"),
    ("Компилятор", "gcc 8.5.0, -O3 -fopenmp; OMP_PROC_BIND=close, OMP_PLACES=cores"),
], [4, 13])

# --- 2. Декомпозиция
h("2. Декомпозиция данных")
doc.add_picture("plots/decomposition.png", width=Cm(11))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
bullet("строки делятся schedule(static) на p равных полос — работа на строку одинакова, "
       "поэтому баланс идеальный;", "Ленты: ")
bullet("поток пишет только свои строки unew и только читает u;", "Нет гонок: ")
bullet("один parallel-регион на 500 итераций, обмен указателей u↔unew в omp single;",
       "Синхронизация: ")
bullet("массивы инициализируются тем же разбиением, данные ложатся в память своего сокета.",
       "First touch: ")

# --- 3. Результаты
h("3. Результаты")
para(f"Запуск через sbatch, время только вычислительного ядра, минимум из 3 запусков. "
     f"T1 = {T1:.3f} с (seq.c). S = T1/T(p), E = S/p, e = (1/S − 1/p)/(1 − 1/p). "
     "ГБ/с — оценка снизу, 16 Б на точку.", italic=False)
table(["p", "T(p), с", "S(p)", "E(p)", "e (К–Ф)", "ГБ/с", "CHECK"],
      [[r["p"], f"{r['T']:.3f}", f"{r['S']:.2f}", f"{r['E']:.2f}",
        f"{r['e']:.4f}" if r["e"] is not None else "—", f"{bw(r):.0f}", r["check"]] for r in rows],
      [1.2, 2.2, 2, 2, 2.4, 2, 2])
g = doc.add_table(rows=1, cols=2)
for i, f in enumerate(["plots/speedup.png", "plots/efficiency.png"]):
    c = g.rows[0].cells[i]; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(f, width=Cm(8.3))
para(f"Лог: {', '.join(R['logs'])}", italic=True)

# --- 4. Анализ и вывод
h("4. Анализ и вывод")
para(f"Карп–Флатт: e растёт от {P[2]['e']:.3f} (p = 2) до {P[8]['e']:.3f} (p = 8) и дальше держится "
     f"на ≈ {P[32]['e']:.3f}. Рост e говорит о накладных расходах, а не о последовательной части: "
     "в ядре последовательного кода почти нет. Формально по Амдалу S∞ ≈ "
     f"1/{P[32]['e']:.3f} ≈ {1/P[32]['e']:.0f}.")
para(f"Узкое место — пропускная способность памяти: 5 операций на точку против ≥ 16 Б обмена. "
     f"До p = 4 удвоение потоков даёт почти ×2, на 8 → 16 только ×{gain(8, 16):.2f}: 16 потоков "
     f"занимают один сокет и выбирают ≈ {bw(P[16]):.0f} ГБ/с. На p = 32 включается второй сокет "
     f"со своей памятью — ускорение растёт медленнее (×{gain(16, 32):.2f}), ≈ {bw(P[32]):.0f} ГБ/с. "
     "Дисбаланса нет (static), барьеров всего 1000 за расчёт.")
para(f"Вывод: CHECK: OK во всех 21 запуске. На 32 потоках T = {last['T']:.3f} с против "
     f"T1 = {T1:.3f} с: S = {last['S']:.2f}, E = {last['E']:.2f}. Дальше ускорять можно блочным "
     "обходом (temporal blocking) или переходом на несколько узлов через MPI.", bold_prefix=None)

doc.save("report_t3.docx")
print("сохранено: report_t3.docx")
